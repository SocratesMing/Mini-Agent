import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class FileParser:
    """文件内容解析器，支持多种文件格式"""

    @staticmethod
    def extract_content(file_path: str, file_type: str = "") -> str:
        """从文件中提取文本内容

        Args:
            file_path: 文件路径
            file_type: 文件类型

        Returns:
            提取的文本内容
        """
        file_path_obj = Path(file_path)
        suffix = file_type.lower() or file_path_obj.suffix.lstrip('.').lower()
        if suffix.startswith('.'):
            suffix = suffix[1:]
        logger.info(f"[FileParser] extract_content 被调用 | file_path: {file_path} | suffix: {suffix}")

        try:
            if suffix == 'pdf':
                return FileParser._extract_pdf(file_path)
            elif suffix in ('doc', 'docx'):
                return FileParser._extract_docx(file_path)
            elif suffix in ('xls', 'xlsx'):
                return FileParser._extract_excel(file_path)
            elif suffix in ('ppt', 'pptx'):
                return FileParser._extract_pptx(file_path)
            elif suffix == 'csv':
                return FileParser._extract_csv(file_path)
            elif suffix in ('json', 'jsonl'):
                return FileParser._extract_json(file_path)
            elif suffix == 'xml':
                return FileParser._extract_xml(file_path)
            elif suffix in ('html', 'htm'):
                return FileParser._extract_html(file_path)
            elif suffix == 'txt':
                return FileParser._extract_txt(file_path)
            elif suffix == 'md':
                return FileParser._extract_txt(file_path)
            else:
                return FileParser._extract_code_file(file_path, suffix)

        except Exception as e:
            logger.error(f"提取文件内容失败：{e}")
            return ""

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        """提取 PDF 内容"""
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(file_path)
            texts = []

            for page in reader.pages:
                text = page.extract_text()
                if text:
                    texts.append(text)

            return "\n\n".join(texts)

        except ImportError:
            logger.error("请安装 PyPDF2: pip install PyPDF2")
            return ""
        except Exception as e:
            logger.error(f"PDF 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        """提取 Word 文档内容"""
        try:
            from docx import Document

            logger.info(f"[FileParser] 尝试打开 docx 文件：{file_path}")
            doc = Document(file_path)
            logger.info(f"[FileParser] docx 文件已打开，段落数：{len(doc.paragraphs)}, 表格数：{len(doc.tables)}")
            texts = []

            for para in doc.paragraphs:
                if para.text.strip():
                    texts.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    row_texts = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_texts.append(cell.text.strip())
                    if row_texts:
                        texts.append(" | ".join(row_texts))

            result = "\n\n".join(texts)
            logger.info(f"[FileParser] docx 提取完成，文本长度：{len(result)}")
            return result

        except ImportError:
            logger.error("请安装 python-docx: pip install python-docx")
            return ""
        except Exception as e:
            logger.error(f"DOCX 提取失败：{e}", exc_info=True)
            return ""

    @staticmethod
    def _extract_excel(file_path: str) -> str:
        """提取 Excel 内容（增强版：支持计息表格）"""
        try:
            import openpyxl
            logger.info(f"[FileParser] 尝试打开 xlsx 文件：{file_path}")

            wb = openpyxl.load_workbook(file_path, data_only=True)
            logger.info(f"[FileParser] xlsx 文件已打开，sheet 数：{len(wb.sheetnames)}")
            texts = []

            # 计息相关的列名关键词
            interest_keywords = ['利息', '利率', '计息', '本金', '金额', '天数', '起息', '到期', '逾期', '罚息', '复利']
            
            # 需要格式化的数值列关键词
            amount_keywords = ['金额', '本金', '利息', '余额', '总额', '合计', '利率', '比例']

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                texts.append(f"[Sheet: {sheet_name}]")

                # 获取表头（第一行）
                headers = []
                header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                if header_row:
                    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
                
                # 检测是否包含计息相关列
                is_interest_sheet = any(
                    any(kw in header for kw in interest_keywords)
                    for header in headers if header
                )
                
                logger.info(f"[FileParser] Sheet '{sheet_name}' 是否计息表：{is_interest_sheet}")
                
                # 识别金额列的索引
                amount_columns = []
                if headers:
                    for i, header in enumerate(headers):
                        if any(kw in header for kw in amount_keywords):
                            amount_columns.append(i)
                
                # 处理数据行（从第 2 行开始）
                for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(cell is not None for cell in row):
                        continue
                    
                    row_texts = []
                    for col_idx, cell in enumerate(row):
                        if cell is None:
                            continue
                        
                        # 如果是金额列，格式化数值
                        if col_idx in amount_columns:
                            try:
                                if isinstance(cell, (int, float)):
                                    header = headers[col_idx] if col_idx < len(headers) else ""
                                    if '利率' in header or '比例' in header:
                                        # 格式化为百分比
                                        formatted = f"{cell * 100:.2f}%"
                                    else:
                                        # 格式化为金额，保留两位小数
                                        formatted = f"{cell:,.2f}"
                                    row_texts.append(f"{headers[col_idx] if col_idx < len(headers) else ''}:{formatted}")
                                    continue
                            except (ValueError, TypeError):
                                pass
                        
                        # 默认处理
                        cell_str = str(cell).strip()
                        if cell_str:
                            if col_idx < len(headers) and headers[col_idx]:
                                row_texts.append(f"{headers[col_idx]}:{cell_str}")
                            else:
                                row_texts.append(cell_str)
                    
                    if row_texts:
                        texts.append(" | ".join(row_texts))

                texts.append("")

            result = "\n".join(texts)
            logger.info(f"[FileParser] xlsx 提取完成，文本长度：{len(result)}")
            return result

        except ImportError:
            logger.error("请安装 openpyxl: pip install openpyxl")
            return ""
        except Exception as e:
            logger.error(f"Excel 提取失败：{e}", exc_info=True)
            return ""

    @staticmethod
    def _extract_txt(file_path: str) -> str:
        """提取文本文件内容"""
        try:
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']

            for encoding in encodings:
                try:
                    content = Path(file_path).read_text(encoding=encoding)
                    logger.info(f"[FileParser] _extract_txt 读取文件：{file_path}, 长度：{len(content)}, 前 100 字符：{repr(content[:100])}")
                    return content
                except UnicodeDecodeError:
                    continue

            return ""

        except Exception as e:
            logger.error(f"文本文件读取失败：{e}")
            return ""

    @staticmethod
    def _extract_pptx(file_path: str) -> str:
        """提取 PowerPoint 内容"""
        try:
            from pptx import Presentation
            import json

            prs = Presentation(file_path)
            slides = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text)
                slides.append({
                    "number": slide_num,
                    "content": "\n".join(slide_texts) if slide_texts else "(空白幻灯片)"
                })

            return json.dumps({"slides": slides}, ensure_ascii=False)

        except ImportError:
            logger.error("请安装 python-pptx: pip install python-pptx")
            return ""
        except Exception as e:
            logger.error(f"PPTX 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_csv(file_path: str) -> str:
        """提取 CSV 内容"""
        try:
            import csv

            texts = []
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin1']

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        reader = csv.reader(f)
                        for row in reader:
                            row_text = " | ".join(str(cell) for cell in row if cell)
                            if row_text:
                                texts.append(row_text)
                    break
                except UnicodeDecodeError:
                    continue

            return "\n".join(texts)

        except Exception as e:
            logger.error(f"CSV 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_json(file_path: str) -> str:
        """提取 JSON 内容"""
        try:
            import json

            encodings = ['utf-8', 'gbk', 'gb2312']
            content = ""

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            if not content:
                return ""

            data = json.loads(content)

            def flatten(obj, prefix=""):
                result = []
                if isinstance(obj, dict):
                    for key, value in obj.items():
                        result.extend(flatten(value, f"{prefix}{key}."))
                elif isinstance(obj, list):
                    for i, item in enumerate(obj):
                        result.extend(flatten(item, f"{prefix}[{i}]."))
                else:
                    result.append(f"{prefix[:-1]}: {obj}")
                return result

            return "\n".join(flatten(data))

        except Exception as e:
            logger.error(f"JSON 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_xml(file_path: str) -> str:
        """提取 XML 内容"""
        try:
            import xml.etree.ElementTree as ET

            encodings = ['utf-8', 'gbk', 'gb2312']
            content = ""

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            if not content:
                return ""

            root = ET.fromstring(content)

            def extract_text(element, prefix=""):
                result = []
                if element.text and element.text.strip():
                    result.append(f"{prefix}{element.tag}: {element.text.strip()}")
                for child in element:
                    result.extend(extract_text(child, f"{prefix}{element.tag}."))
                return result

            return "\n".join(extract_text(root))

        except Exception as e:
            logger.error(f"XML 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_html(file_path: str) -> str:
        """提取 HTML 内容"""
        try:
            from bs4 import BeautifulSoup

            encodings = ['utf-8', 'gbk', 'gb2312']
            content = ""

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            if not content:
                return ""

            soup = BeautifulSoup(content, 'html.parser')

            for script in soup(["script", "style"]):
                script.decompose()

            text = soup.get_text(separator='\n', strip=True)

            lines = [line for line in text if line]
            return "\n".join(lines)

        except ImportError:
            logger.error("请安装 beautifulsoup4: pip install beautifulsoup4")
            return ""
        except Exception as e:
            logger.error(f"HTML 提取失败：{e}")
            return ""

    @staticmethod
    def _extract_code_file(file_path: str, suffix: str) -> str:
        """提取代码文件内容"""
        try:
            encodings = ['utf-8', 'gbk', 'gb2312']
            content = ""

            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            if not content:
                return ""

            code_extensions = {
                'py': 'Python', 'js': 'JavaScript', 'ts': 'TypeScript', 'jsx': 'React JSX',
                'tsx': 'React TSX', 'java': 'Java', 'c': 'C', 'cpp': 'C++', 'h': 'C Header',
                'hpp': 'C++ Header', 'cs': 'C#', 'go': 'Go', 'rs': 'Rust', 'php': 'PHP',
                'rb': 'Ruby', 'swift': 'Swift', 'kt': 'Kotlin', 'scala': 'Scala',
                'lua': 'Lua', 'sh': 'Shell', 'bash': 'Bash', 'zsh': 'Zsh',
                'ps1': 'PowerShell', 'sql': 'SQL', 'r': 'R', 'm': 'MATLAB/Objective-C',
                'vb': 'Visual Basic', 'pl': 'Perl', 'hs': 'Haskell', 'ex': 'Elixir',
                'exs': 'Elixir', 'erl': 'Erlang', 'fs': 'F#', 'clj': 'Clojure',
                'dart': 'Dart', 'groovy': 'Groovy', 'gradle': 'Gradle', 'toml': 'TOML',
                'yaml': 'YAML', 'yml': 'YAML', 'json': 'JSON', 'xml': 'XML',
                'html': 'HTML', 'css': 'CSS', 'scss': 'SCSS', 'sass': 'Sass',
                'less': 'Less', 'vue': 'Vue', 'svelte': 'Svelte', 'md': 'Markdown',
                'rst': 'reStructuredText', 'tex': 'LaTeX', 'ini': 'INI', 'cfg': 'Config',
                'conf': 'Config', 'env': 'Environment', 'properties': 'Properties',
                'gitignore': 'Git Ignore', 'dockerfile': 'Dockerfile', 'makefile': 'Makefile',
            }

            lang = code_extensions.get(suffix.lower(), suffix.upper())
            lines = content.split('\n')
            code_lines = []
            for i, line in enumerate(lines, 1):
                code_lines.append(f"{i:4d} | {line}")

            return f"// {lang} Code File\n// Lines: {len(lines)}\n\n" + "\n".join(code_lines)

        except Exception as e:
            logger.error(f"代码文件读取失败：{e}")
            return ""
